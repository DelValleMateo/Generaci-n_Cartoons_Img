"""Construye el informe Word a partir de los CSV reales ya validados."""
import csv
from collections import Counter
from pathlib import Path
from statistics import mean
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from validar_resultados import check

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'resultados' / 'matriz_completa'
OUT = ROOT / 'documentacion' / 'informe_final.docx'
check()
with (DATA / 'resumen.csv').open(newline='', encoding='utf-8') as source:
    ROWS = list(csv.DictReader(source))
CASES = sorted({(r['tamano'], int(r['filtro']), int(r['niveles'])) for r in ROWS},
               key=lambda k: (int(k[0].split('x')[0]), k[1], k[2]))
BY_CASE = {case: [r for r in ROWS if (r['tamano'], int(r['filtro']), int(r['niveles'])) == case]
           for case in CASES}
WINNERS = {case: min(items, key=lambda r: float(r['tiempo_promedio_s'])) for case, items in BY_CASE.items()}
COUNTS = Counter(r['implementacion'] for r in WINNERS.values())
LABEL = {'secuencial': 'Secuencial', 'openmp': 'OpenMP', 'mpi': 'MPI', 'hibrido': 'Híbrida'}


def get(size, filt, levels, implementation, processes, threads):
    return next(r for r in ROWS if r['tamano'] == f'{size}x{size}' and int(r['filtro']) == filt
                and int(r['niveles']) == levels and r['implementacion'] == implementation
                and int(r['procesos']) == processes and int(r['hilos']) == threads)


def average_speedup(size, implementation, processes, threads):
    return mean(float(get(size, f, l, implementation, processes, threads)['speedup'])
                for f in (3, 5) for l in (3, 9))


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.4)
section.bottom_margin = Cm(2.3)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Title'].font.size = Pt(17)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(12.5)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.keep_with_next = True
doc.styles['Caption'].font.size = Pt(9)


def paragraph(text):
    return doc.add_paragraph(text)


def heading(number, title, text):
    doc.add_heading(f'{number} {title}', level=1)
    paragraph(text)


def table(headers, rows, widths=None):
    item = doc.add_table(rows=1, cols=len(headers))
    item.style = 'Table Grid'
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    for i, value in enumerate(headers):
        cell = item.rows[0].cells[i]
        cell.text = str(value)
        if widths:
            cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = 'Arial'
            run.font.size = Pt(7.5)
    header = item.rows[0]._tr
    tr_pr = header.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)
    for values in rows:
        cells = item.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if widths:
                cells[i].width = Cm(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(7.5)
    doc.add_paragraph()
    return item


def figure(filename, caption):
    path = ROOT / 'graficos' / filename
    assert path.exists(), path
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(path), width=Cm(15.7))
    p = doc.add_paragraph(caption, style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc.add_paragraph('Generación de imágenes cartoon con MPI y OpenMP', style='Title')
paragraph('Práctico Final Integrador de Computación Avanzada · UADER · 2026')
paragraph('Resultados medidos el 28 de octubre de 2026 en una computadora Windows')
heading('1', 'Introducción y objetivos',
        'El objetivo fue construir una imagen cartoon a partir de una imagen RGB y comparar una implementación secuencial con tres formas de paralelismo: MPI, OpenMP y una combinación de ambas. El análisis utiliza exclusivamente las ejecuciones conservadas en los CSV del proyecto. Cada salida paralela se comparó con la secuencial para el mismo tamaño, filtro y cantidad de niveles; la equivalencia binaria se confirmó en los doce casos exigidos por la consigna (UADER, 2026).')
heading('2', 'Fundamentos del procesamiento de imágenes',
        'Una imagen RGB de 8 bits almacena tres canales por píxel, cada uno entre 0 y 255. Para detectar contornos se calcula una intensidad gris como 0,30R + 0,59G + 0,11B. El promedio espacial de una vecindad atenúa variaciones locales antes de aplicar Sobel, que estima cambios horizontales y verticales. La posterización reduce el número de valores posibles por canal y simplifica las regiones de color (Guerrero, 2017).')
heading('3', 'Algoritmo cartoon',
        'El programa convierte la imagen a gris, aplica un promedio uniforme de 3 × 3 o 5 × 5 píxeles, calcula |Sx| + |Sy| con Sobel y marca un borde cuando el resultado supera el umbral 150. En paralelo lógico, cada canal RGB se cuantiza en 3 o 9 franjas, usando el punto medio de cada franja. La salida suma la máscara de bordes a la imagen posterizada y limita cada canal a 255. La opción de bordes oscuros resta esa máscara y limita inferiormente a cero; los benchmarks emplearon bordes claros.')
heading('4', 'Implementación secuencial',
        'La referencia recorre todos los píxeles en un proceso. Las etapas producen vectores separados de gris, imagen borrosa y bordes; luego se posteriza una copia RGB y se combinan los resultados. El cronómetro steady_clock comienza después de leer el PPM y termina antes de escribir la salida. Esta medición define la base para cada speedup.')
heading('5', 'Implementación MPI',
        'El proceso raíz lee el PPM y difunde ancho y alto con MPI_Bcast. Las filas se reparten en bloques contiguos, asignando las filas sobrantes a los primeros procesos. MPI_Scatterv envía bloques de tamaño variable con halos, y MPI_Gatherv reconstruye la imagen con las filas propias de cada proceso. La medición con MPI_Wtime comienza antes del reparto y termina después de la recolección, por lo que incluye esas comunicaciones. El código informa error si la lectura, la configuración o la escritura falla.')
heading('6', 'Implementación OpenMP',
        'Las etapas de gris, borroneo, Sobel, posterización y combinación emplean parallel for. Cada iteración escribe una posición distinta; los acumuladores del promedio son variables locales del bucle. La barrera implícita al finalizar cada región asegura que la etapa siguiente lea una matriz completa. Se deshabilitó el ajuste dinámico de hilos y se informó el número efectivo. Se comprobaron 1, 2, 4 y 8 hilos.')
heading('7', 'Implementación híbrida MPI y OpenMP',
        'Cada proceso MPI recibe un bloque con halos y ejecuta las etapas de píxeles con hilos OpenMP. Las llamadas MPI quedan fuera de las regiones paralelas. Se midieron 2 procesos × 2 hilos, 2 × 4 y 4 × 2. El programa informa el mínimo y máximo de hilos efectivos observados entre procesos; la prueba de casos límite verificó que coincidieran con la configuración solicitada.')
heading('8', 'División del problema',
        'El borroneo de tamaño k necesita k/2 filas vecinas a cada lado, mientras que Sobel necesita una fila más sobre el resultado borroso. Por ello se distribuye un halo de k/2 + 1 filas: dos para 3 × 3 y tres para 5 × 5. Solo las filas propias vuelven al proceso raíz. El reparto contiguo simplifica las transferencias, aunque agrega cálculo redundante en los halos y copias de memoria.')
heading('9', 'Comunicación y sincronización',
        'MPI_Barrier alinea el inicio de la región cronometrada y las operaciones colectivas coordinan reparto y recolección. No hay intercambio entre procesos durante el cálculo porque cada uno recibió el halo completo. Las barreras implícitas de OpenMP separan etapas dependientes. La comparación byte por byte de todas las salidas frente a la referencia ayuda a detectar errores en fronteras de bloques y carreras visibles en el resultado.')
heading('10', 'Metodología y equipo de pruebas',
        'Se generaron escenas RGB deterministas de 800 × 800, 2000 × 2000 y 5000 × 5000 píxeles. Para cada tamaño se probaron filtros de 3 × 3 y 5 × 5 y posterización de 3 y 9 niveles: doce casos. Cada caso se ejecutó tres veces en once configuraciones, para un total de 396 mediciones. Los tiempos son los impresos por cada programa y excluyen lectura y escritura. El equipo utilizado tiene un AMD Ryzen 7 5700 y 16 procesadores lógicos; ejecutó Windows build 26200, GCC 16.2 y Microsoft MPI 10.1.3. Las pruebas se hicieron en un solo equipo, de modo que MPI mide comunicación entre procesos locales. La auditoría de CSV verificó tres repeticiones únicas por configuración, tiempos positivos, promedios, fórmulas y un único SHA-256 de salida por cada uno de los doce casos.')
heading('11', 'Resultados de rendimiento',
        f'OpenMP con 8 hilos obtuvo el menor promedio en {COUNTS["openmp"]} de los doce casos; la configuración híbrida 4 × 2 lo obtuvo en {COUNTS["hibrido"]}. La tabla 1 muestra la configuración ganadora de cada caso. El archivo resumen.csv conserva las 132 filas con procesos, hilos, tiempo, speedup y eficiencia; mediciones.csv conserva las 396 observaciones individuales.')
winner_rows = []
for case in CASES:
    r = WINNERS[case]
    winner_rows.append((r['tamano'], f"{case[1]}×{case[1]}", case[2], LABEL[r['implementacion']],
                        r['procesos'], r['hilos'], f"{float(r['tiempo_promedio_s']):.4f}",
                        f"{float(r['speedup']):.2f}", f"{float(r['eficiencia']):.2f}"))
table(('Imagen', 'Filtro', 'Niv.', 'Versión', 'Proc.', 'Hilos', 'Tiempo s', 'Speedup', 'Efic.'),
      winner_rows, (2.0, 1.3, 1.0, 1.9, 1.1, 1.1, 2.0, 1.7, 1.4))
paragraph('Tabla 1. Configuración con menor tiempo promedio en cada caso. Speedup y eficiencia corresponden a esa misma configuración.')
figure('tiempos_comparativos.png', 'Figura 1. Tiempo de las cuatro implementaciones en configuraciones representativas.')
figure('coste_filtro.png', 'Figura 2. Relación entre los tiempos de filtros 5 × 5 y 3 × 3.')
paragraph('El filtro de 5 × 5 exige hasta 25 contribuciones por píxel, frente a 9 del filtro 3 × 3. En la versión secuencial, la relación media tiempo 5 × 5 / tiempo 3 × 3, promediada sobre los dos niveles, fue 1,45 en 800 × 800, 1,59 en 2000 × 2000 y 1,61 en 5000 × 5000. En algunas configuraciones paralelas pequeñas la relación cambia por variación de ejecución y sobrecarga, por lo que el costo del filtro no debe inferirse de un solo punto.')
heading('12', 'Speedup y eficiencia',
        'Para cada caso se calculó speedup = promedio secuencial / promedio paralelo de la misma imagen, filtro y niveles. La eficiencia divide ese speedup por hilos para OpenMP, procesos para MPI o procesos × hilos para la híbrida. La tabla 2 resume la media de los cuatro speedups de filtro y niveles para cada tamaño. Son medias descriptivas de razones por caso, no el cociente entre tiempos agregados.')
table(('Imagen', 'OpenMP 8 h', 'MPI 4 p', 'Híbrida 2×4', 'Híbrida 4×2'),
      [(f'{size}×{size}', f'{average_speedup(size,"openmp",1,8):.2f}',
        f'{average_speedup(size,"mpi",4,1):.2f}',
        f'{average_speedup(size,"hibrido",2,4):.2f}',
        f'{average_speedup(size,"hibrido",4,2):.2f}') for size in (800, 2000, 5000)],
      (2.6, 2.9, 2.9, 2.9, 2.9))
paragraph('Tabla 2. Speedup medio (veces) sobre las cuatro combinaciones de filtro y niveles de cada tamaño.')
figure('speedup_openmp.png', 'Figura 3. Escalamiento de OpenMP entre 1 y 8 hilos.')
figure('speedup_mpi.png', 'Figura 4. Escalamiento de MPI entre 1 y 4 procesos.')
figure('speedup_hibrido.png', 'Figura 5. Comparación de tres configuraciones híbridas.')
figure('eficiencia.png', 'Figura 6. Eficiencia de configuraciones paralelas representativas.')
paragraph('En 5000 × 5000, el speedup medio de OpenMP con 8 hilos fue 3,51 y el de MPI con 4 procesos fue 2,38. Las dos configuraciones híbridas de ocho unidades obtuvieron 2,94 (2 × 4) y 3,16 (4 × 2). La eficiencia correspondiente a OpenMP de 8 hilos fue aproximadamente 0,44, frente a 0,39 para la híbrida 4 × 2, al promediar los cuatro casos. El aumento de unidades mejora el tiempo, pero con rendimiento marginal decreciente.')
heading('13', 'Comparación de implementaciones',
        'OpenMP fue la opción más rápida en la mayoría de los casos medidos en esta máquina compartida. MPI escaló al pasar de uno a cuatro procesos, aunque su ejecución de un solo proceso fue más lenta que la referencia; además del cómputo, su tiempo incluye Scatterv, Gatherv y las copias de bloques. Este resultado es compatible con un costo de comunicación, pero el ensayo no separa ese costo de las otras diferencias de implementación. La híbrida 4 × 2 ganó cuatro casos y superó a la 2 × 4 en la media de speedup de 5000 × 5000. No se extrapolan estas comparaciones a un clúster ni a otros procesadores.')
paragraph('El promedio de speedup de OpenMP con 8 hilos creció de 2,38 en 800 × 800 a 2,88 en 2000 × 2000 y 3,51 en 5000 × 5000. El de MPI con 4 procesos pasó de 2,10 a 2,28 y 2,38. El aumento del trabajo por ejecución amortiza una parte de los costos fijos de iniciar regiones paralelas y mover datos, aunque el código no mide cada componente por separado. Con un único hilo, OpenMP alcanzó speedups medios de 0,81, 0,80 y 0,84 según el tamaño; con un solo proceso, MPI obtuvo 0,77, 0,78 y 0,83. Esas variantes ejecutan trabajo adicional y por ello no son equivalentes en costo a la referencia secuencial.')
paragraph('La distribución MPI de este programa se hace desde el proceso raíz mediante Scatterv de bloques ya ampliados; no se realiza un intercambio directo de halos entre vecinos durante el cálculo. Esta decisión simplifica el control de fronteras, pero agrega bytes enviados y trabajo redundante a medida que crece el número de procesos. La alternativa híbrida reduce el tamaño del bloque por proceso y lo divide entre hilos. En los cuatro casos de 5000 × 5000, la media de speedup fue 2,94 para 2 × 4 y 3,16 para 4 × 2; la diferencia describe este ensayo y no demuestra que 4 × 2 sea siempre superior.')
heading('14', 'Ventajas desventajas y limitaciones',
        'La secuencial ofrece una referencia simple. OpenMP evita repartir la imagen entre espacios de memoria y requiere cuidar la independencia de las escrituras y las barreras entre etapas. MPI permite pensar el algoritmo para memoria distribuida, a cambio de halos y comunicaciones. La híbrida combina ambos esquemas, pero agrega sincronización y copias; una mayor cantidad de unidades no garantiza eficiencia alta. Se usó PPM P6 para evitar dependencias de decodificación en C++, con mayor tamaño de archivo que JPEG o PNG. Las entradas de benchmark son sintéticas y deterministas; una fotografía real de 1330 × 1002 se procesó aparte y sus cuatro salidas coincidieron en SHA-256. Las mediciones de tiempos cortos pueden fluctuar por la resolución del reloj y la actividad del sistema; hay tres repeticiones, sin intervalos de confianza.')
paragraph('El alcance experimental está limitado a un equipo Windows y a una implementación concreta de cada modelo. La medida secuencial y OpenMP abarca solo cómputo, mientras que MPI e híbrida incluyen distribución y reunión de datos; por ello el speedup compara el tiempo observado de cada programa, no kernels de cálculo idénticos. No se separaron tiempos de comunicación, creación de hilos, asignación de memoria ni entrada y salida. La repetición de tres corridas permite calcular medias y observar dispersión individual en mediciones.csv, pero no permite afirmar diferencias pequeñas con significación estadística.')
heading('15', 'Conclusiones',
        'Las cuatro implementaciones compilaron y produjeron resultados equivalentes en toda la matriz exigida. Las pruebas adicionales compararon 260 salidas pequeñas píxel por píxel con una referencia independiente y comprobaron 32 rechazos de entradas, parámetros o destinos inválidos. En el hardware utilizado, OpenMP con 8 hilos lideró ocho casos y la híbrida 4 × 2 lideró cuatro. El rendimiento mejoró con más trabajo por imagen, pero la eficiencia permaneció por debajo de uno al aumentar las unidades. La distribución por filas con halos resolvió correctamente las dependencias espaciales sin artefactos en las fronteras de proceso.')
heading('16', 'Referencias bibliográficas',
        'Guerrero, R. (2017). Procesamiento de imágenes [Apuntes de cátedra]. Archivo Apuntes.pdf.\nUADER. (2026). Computación Avanzada: Práctico Final Integrador [Consigna de cátedra]. Archivo CA-Practico Final-2026.pdf.')
paragraph('Datos de respaldo: resultados/matriz_completa/mediciones.csv, resumen.csv y validacion.txt. Los seis gráficos se generaron exclusivamente desde resumen.csv.')
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
